"""`@component:` directives — the mechanism that lets a .docx name an MDX
component without the pipeline knowing what that component is.

Every failure mode here degrades to the plain block plus a warning: an editor
mistyping a component name must not be able to break a page.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from drive_sync import mapping as mapping_mod
from drive_sync.emit.format import emit_blocks
from drive_sync.mapping import load_components
from drive_sync.models import Component, Table
from drive_sync.parse.docx import ParseDocxOptions, parse_docx
from drive_sync.report import ParseReport


def _doc_with(tmp_path: Path, directive: str | None, rows: list[list[str]]) -> Path:
    """A minimal valid info.docx whose Introduction holds one optional
    directive followed by a table."""
    doc = Document()
    doc.add_heading("Metadata", level=1)
    meta = doc.add_table(rows=4, cols=2)
    for r, (k, v) in zip(
        meta.rows,
        [("Key", "Value"), ("title", "T"), ("sidebar_label", "T"), ("sidebar_position", "1")],
    ):
        r.cells[0].text = k
        r.cells[1].text = v

    doc.add_heading("Introduction", level=1)
    if directive:
        doc.add_paragraph(directive)
    if rows:
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, values in zip(t.rows, rows):
            for cell, value in zip(r.cells, values):
                cell.text = value
    path = tmp_path / "info.docx"
    doc.save(str(path))
    return path


def _parse(path: Path):
    report = ParseReport()
    parsed = parse_docx(str(path), ParseDocxOptions(file_label="t/info.docx"), report)
    assert parsed is not None, [e.message for e in report.entries]
    return parsed.sections["Introduction"], report


TUITION = [["Faculty", "USD per credit", "Reference"], ["MSFEA", "990", "https://x"]]


def test_directive_turns_the_next_table_into_a_component(tmp_path: Path) -> None:
    blocks, report = _parse(_doc_with(tmp_path, "@component: TuitionTable", TUITION))
    comps = [b for b in blocks if isinstance(b, Component)]
    assert len(comps) == 1
    c = comps[0]
    assert c.name == "TuitionTable"
    assert c.rows == [{"faculty": "MSFEA", "usd_per_credit": "990", "reference": "https://x"}]
    assert not report.has_errors()
    assert not any("@component" in str(b) for b in blocks)


def test_component_emits_jsx_with_numeric_columns_unquoted(tmp_path: Path) -> None:
    blocks, _ = _parse(_doc_with(tmp_path, "@component: TuitionTable", TUITION))
    out = emit_blocks(blocks)
    assert "<TuitionTable" in out
    assert "faculty: 'MSFEA'" in out
    assert "usd_per_credit: 990" in out, "declared numeric columns are not strings"


def test_directive_args_become_props(tmp_path: Path) -> None:
    blocks, _ = _parse(
        _doc_with(tmp_path, '@component: TuitionTable year="2025-2026"', TUITION)
    )
    c = [b for b in blocks if isinstance(b, Component)][0]
    assert c.props == {"year": "2025-2026"}
    assert "year='2025-2026'" in emit_blocks(blocks)


def test_table_without_a_directive_is_untouched(tmp_path: Path) -> None:
    """The default path stays a plain markdown table — directives are opt-in."""
    blocks, report = _parse(_doc_with(tmp_path, None, TUITION))
    assert any(isinstance(b, Table) for b in blocks)
    assert not any(isinstance(b, Component) for b in blocks)
    assert report.count("warning") == 0


def test_unknown_component_falls_back_to_a_plain_table_and_warns(tmp_path: Path) -> None:
    blocks, report = _parse(_doc_with(tmp_path, "@component: TuitionTabel", TUITION))
    assert any(isinstance(b, Table) for b in blocks), "page still renders"
    assert not any(isinstance(b, Component) for b in blocks)
    assert not report.has_errors(), "a typo must not fail the build"
    assert any("unknown component" in e.message for e in report.entries)


def test_missing_required_column_falls_back_and_warns(tmp_path: Path) -> None:
    """MajorsTable requires `program`; a table without it is clearly not one."""
    blocks, report = _parse(
        _doc_with(tmp_path, "@component: MajorsTable", [["School", "Rate"], ["MSFEA", "990"]])
    )
    assert any(isinstance(b, Table) for b in blocks)
    assert not report.has_errors()
    assert any("missing required column" in e.message for e in report.entries)


def test_children_source_wraps_the_block(tmp_path: Path) -> None:
    doc_path = _doc_with(tmp_path, None, [])
    doc = Document(str(doc_path))
    doc.add_paragraph("@component: Callout variant=warning")
    doc.add_paragraph("Deadlines moved this year.")
    doc.save(str(doc_path))

    blocks, report = _parse(doc_path)
    comps = [b for b in blocks if isinstance(b, Component)]
    assert len(comps) == 1 and comps[0].name == "Callout"
    assert comps[0].children, "the paragraph below became the children"
    out = emit_blocks(blocks)
    assert "<Callout variant='warning'>" in out and "</Callout>" in out
    assert not report.has_errors()


def test_a_new_component_needs_no_python_change(tmp_path: Path, monkeypatch) -> None:
    """The acceptance test for the directive design.

    A component the codebase has never heard of becomes usable by adding one
    row to components.toml. No parser, emitter or model change.
    """
    registry = tmp_path / "components.toml"
    registry.write_text(
        """
[RankingsTable]
source = "table"
required = ["ranking"]
numeric = ["position"]
""",
        encoding="utf-8",
    )
    monkeypatch.setattr(mapping_mod, "COMPONENTS_PATH", registry)
    load_components.cache_clear()
    try:
        blocks, report = _parse(
            _doc_with(
                tmp_path,
                "@component: RankingsTable",
                [["Ranking", "Position"], ["QS World", "601"]],
            )
        )
        c = [b for b in blocks if isinstance(b, Component)][0]
        assert c.name == "RankingsTable"
        assert c.rows == [{"ranking": "QS World", "position": "601"}]
        assert "position: 601" in emit_blocks(blocks)
        assert not report.has_errors()
    finally:
        load_components.cache_clear()


def test_every_registered_component_is_available_to_mdx() -> None:
    """components.toml and MDXComponents.tsx must not drift.

    The pipeline will happily emit `<TuitionTable>`; if the site has not
    registered that name, MDX renders it as an unknown element and the build
    fails. This is the only thing holding the two registries together.
    """
    repo = Path(__file__).resolve().parents[3]
    registered = (repo / "src" / "theme" / "MDXComponents.tsx").read_text(encoding="utf-8")
    missing = [name for name in load_components() if name not in registered]
    assert not missing, (
        f"declared in components.toml but not registered in MDXComponents.tsx: {missing}"
    )
